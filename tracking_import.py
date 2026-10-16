"""Cập nhật bảng Theo dõi ticket (ticket_tracking) từ file Excel log cập nhật hệ thống Gree App.

Đọc sheet "Logs", lấy các dòng có cột "Trạng thái cập nhật" KHÁC trống / "Đã cập nhật" / "Không cập nhật"
(tức là việc còn đang theo dõi), rồi khớp với bảng theo ID AEGONA (dòng không có ID thì khớp theo Tiêu đề):
  - đã có → cập nhật trạng thái/tiêu đề/nội dung... (đổi trạng thái thì tự ghi log cho báo cáo tuần)
  - chưa có → thêm mới (nhóm Gree App)
Mục đang có trong bảng mà file không còn thì giữ nguyên (tự đổi trạng thái trên web nếu cần).

Chạy từ máy (cần .streamlit/secrets.toml để ghi vào DB thật):
    python tracking_import.py <file.xlsx> --dry-run    (chỉ xem thay đổi, không ghi DB)
    python tracking_import.py <file.xlsx>
Trang "Báo cáo tuần" cũng có ô tải file lên dùng chung các hàm này.
"""
import datetime
import sys

SHEET_NAME = "logs"
SKIP_STATUSES = {"", "đã cập nhật", "không cập nhật"}
# Tên cột trong file (so khớp phần đầu, không phân biệt hoa thường) → trường trong ticket_tracking
COLUMN_MAP = {
    "trạng thái cập nhật": "status",
    "ưu tiên": "priority",
    "ngày cập nhật lên server": "prod_date",
    "id aegona": "ext_id",
    "tiêu đề": "title",
    "nội dung cập nhật": "content",
    "source code": "source_code",
    "người thực hiện": "assignee",
}


def _clean(value):
    if value is None:
        return None
    if isinstance(value, datetime.datetime):
        return value.date()
    if isinstance(value, datetime.date):
        return value
    text = str(value).strip()
    return text or None


def normalize_ext_id(value):
    text = str(value or "").strip()
    if not text:
        return None
    return text if text.startswith("#") else f"#{text}"


def parse_tracking_xlsx(file):
    """file: đường dẫn hoặc file-like. Trả về (items, errors) — items: list dict theo các trường COLUMN_MAP."""
    from openpyxl import load_workbook

    wb = load_workbook(file, data_only=True, read_only=True)
    sheet = next((wb[n] for n in wb.sheetnames if n.strip().lower() == SHEET_NAME), None)
    if sheet is None:
        return [], [f"Không có sheet 'Logs' (các sheet: {', '.join(wb.sheetnames)})"]

    rows = sheet.iter_rows(values_only=True)
    header = [str(h or "").strip().lower() for h in next(rows, [])]
    col_idx = {}
    for i, h in enumerate(header):
        for prefix, field in COLUMN_MAP.items():
            if field not in col_idx and h.startswith(prefix):
                col_idx[field] = i
    missing = [p for p, f in COLUMN_MAP.items() if f not in col_idx and f in ("status", "title")]
    if missing:
        return [], [f"Thiếu cột: {', '.join(missing)}"]

    items = []
    for row in rows:
        status = _clean(row[col_idx["status"]]) if col_idx["status"] < len(row) else None
        if str(status or "").lower() in SKIP_STATUSES:
            continue
        item = {f: _clean(row[i]) if i < len(row) else None for f, i in col_idx.items()}
        if not item.get("title"):
            continue
        item["ext_id"] = normalize_ext_id(item.get("ext_id"))
        if item.get("prod_date") is not None and not isinstance(item["prod_date"], datetime.date):
            item["prod_date"] = None
        items.append(item)
    wb.close()
    return items, []


def match_tracking(items, existing):
    """So file với bảng hiện có → (added, updated). updated chỉ chứa id + các trường thay đổi."""
    by_ext = {normalize_ext_id(e.get("ext_id")): e for e in existing if e.get("ext_id")}
    by_title = {str(e.get("title") or "").strip().casefold(): e for e in existing if not e.get("ext_id")}
    added, updated = [], []
    for item in items:
        if item.get("ext_id"):
            current = by_ext.get(item["ext_id"])
        else:
            current = by_title.get(item["title"].strip().casefold())
        if current is None:
            added.append(dict(item, category="gree_app"))
            continue
        changes = {f: v for f, v in item.items() if (current.get(f) or None) != v}
        if changes:
            updated.append(dict(changes, id=current["id"]))
    return added, updated


def main(argv):
    args = [a for a in argv if not a.startswith("--")]
    if not args:
        print(__doc__)
        return 1
    items, errors = parse_tracking_xlsx(args[0])
    if errors:
        print("Lỗi:", *errors, sep="\n- ")
        return 1

    import db
    existing = db.get_ticket_tracking()
    added, updated = match_tracking(items, existing)
    by_id = {e["id"]: e for e in existing}
    print(f"File có {len(items)} mục đang theo dõi → thêm {len(added)}, cập nhật {len(updated)}")
    for item in added:
        print(f"  + {item.get('ext_id') or '(không ID)'} [{item['status']}] {item['title'].splitlines()[0][:70]}")
    for item in updated:
        cur = by_id[item["id"]]
        fields = ", ".join(f"{f}: {cur.get(f)!r:.30} → {v!r:.30}" if f == "status" else f for f, v in item.items() if f != "id")
        print(f"  ~ {cur.get('ext_id') or cur['title'][:40]}: {fields}")
    if "--dry-run" in argv:
        print("(--dry-run: chưa ghi DB)")
        return 0
    if (added or updated) and not db.save_ticket_tracking_changes(added, updated, []):
        print("Không kết nối được DB.")
        return 1
    print("Đã cập nhật DB.")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
